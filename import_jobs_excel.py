"""Import a dated Excel job dataset into the persistent recent-job collection."""
from argparse import ArgumentParser
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from app import EXCEL_COLUMNS, clean_text, load_collected, merge_unique, recent_jobs, save_collected
from collect_jobs import write_xlsx_safely


def read_excel_jobs(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        headers = [clean_text(value).lower().replace(" ", "_") for value in next(iterator)]
        required = {"title", "company", "location", "published_at"}
        missing = required.difference(headers)
        if missing:
            raise ValueError("Missing required columns: " + ", ".join(sorted(missing)))
        rows = []
        for values in iterator:
            source = dict(zip(headers, values))
            if not clean_text(source.get("title")):
                continue
            row = {column: clean_text(source.get(column, "")) for column in EXCEL_COLUMNS}
            row["collected_at"] = row["collected_at"] or datetime.now(timezone.utc).isoformat(timespec="seconds")
            row["verification_status"] = row["verification_status"] or "Imported dataset - link not independently verified"
            row["application_link_or_employer_email"] = row["application_link_or_employer_email"] or row["application_url"] or "Not provided by source"
            rows.append(row)
        return rows
    finally:
        workbook.close()


def main():
    parser = ArgumentParser(description="Append date-eligible Excel rows without duplicates.")
    parser.add_argument("input", help="Input .xlsx file")
    parser.add_argument("--output", default="output/job_role_finder_combined.xlsx")
    args = parser.parse_args()
    imported = read_excel_jobs(Path(args.input))
    eligible = recent_jobs(imported)
    combined, added = merge_unique(load_collected(), eligible)
    save_collected(combined)
    write_xlsx_safely(Path(args.output), combined)
    print(f"Read: {len(imported)} | Eligible: {len(eligible)} | Added: {added} | Final unique: {len(combined)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
